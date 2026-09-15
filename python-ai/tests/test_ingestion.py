import io

import pytest
from docx import Document
from fpdf import FPDF
from pptx import Presentation

from app.ingestion import parse_units


def test_parses_dockerfile_by_name_with_no_extension() -> None:
    content = b"FROM python:3.11\nRUN pip install fastapi\nCMD [\"python\", \"app.py\"]\n"
    units = parse_units("Dockerfile", content)
    assert units
    assert any("FROM python" in unit for unit in units)


def test_parses_terraform_variable_file() -> None:
    content = b'variable "region" {\n  default = "us-east-1"\n}\n\nresource "aws_s3_bucket" "data" {\n  bucket = "example"\n}\n'
    units = parse_units("main.tfvars", content)
    assert any("region" in unit for unit in units)
    assert any("aws_s3_bucket" in unit for unit in units)


def test_parses_ini_style_config_with_sections() -> None:
    content = b"[database]\nhost = localhost\nport = 5432\n\n# comment line\n[cache]\nttl=60\n"
    units = parse_units("service.conf", content)
    assert any("Section: database" in unit and "host = localhost" in unit for unit in units)
    assert any("Section: cache" in unit and "ttl=60" in unit for unit in units)


def test_parses_env_file() -> None:
    content = b"API_KEY=secret\nDEBUG=true\n"
    units = parse_units(".env", content)
    assert any("API_KEY=secret" in unit for unit in units)


def test_parses_html_document_by_block() -> None:
    content = b"<html><body><h1>Title</h1><p>First paragraph.</p><p>Second paragraph.</p></body></html>"
    units = parse_units("page.html", content)
    assert any("Title" in unit for unit in units)
    assert any("First paragraph." in unit for unit in units)
    assert any("Second paragraph." in unit for unit in units)


def test_falls_back_to_generic_xml_when_not_drawio() -> None:
    content = b"<catalog><service><name>Billing</name><owner>Finance</owner></service></catalog>"
    units = parse_units("service-catalog.xml", content)
    assert any("Billing" in unit for unit in units)
    assert any("Finance" in unit for unit in units)


def test_drawio_file_still_uses_diagram_parser() -> None:
    content = (
        b'<mxGraphModel><root>'
        b'<mxCell id="1" vertex="1" value="Order Service"/>'
        b'<mxCell id="2" vertex="1" value="Payment Service"/>'
        b'<mxCell id="3" edge="1" source="1" target="2" value="calls"/>'
        b"</root></mxGraphModel>"
    )
    units = parse_units("flow.drawio", content)
    assert any("Order Service" in unit for unit in units)
    assert any("Relationship: Order Service -> Payment Service" in unit for unit in units)


def test_parses_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Payment Service handles checkout transactions.")
    document.add_paragraph("It depends on Fraud Service for risk checks.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Owner"
    table.rows[0].cells[1].text = "Commerce Platform"
    buffer = io.BytesIO()
    document.save(buffer)

    units = parse_units("architecture-notes.docx", buffer.getvalue())
    assert any("Payment Service handles checkout transactions." in unit for unit in units)
    assert any("Fraud Service" in unit for unit in units)
    assert any("Table row 1" in unit and "Commerce Platform" in unit for unit in units)


def test_parses_pptx_slides_with_notes() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Payment Architecture"
    slide.placeholders[1].text = "Calls Fraud Service\nWrites to Payment PostgreSQL"
    slide.notes_slide.notes_text_frame.text = "Discuss rollout timeline."
    buffer = io.BytesIO()
    presentation.save(buffer)

    units = parse_units("architecture.pptx", buffer.getvalue())
    assert any("Slide 1" in unit and "Payment Architecture" in unit for unit in units)
    assert any("Fraud Service" in unit for unit in units)
    assert any("Speaker notes: Discuss rollout timeline." in unit for unit in units)


def test_parses_pdf_pages() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 10, "Payment Service depends on Fraud Service.")
    content = bytes(pdf.output())

    units = parse_units("architecture-notes.pdf", content)
    assert any("Page 1" in unit and "Payment Service depends on Fraud Service." in unit for unit in units)


def test_pdf_with_no_extractable_text_yields_no_units() -> None:
    pdf = FPDF()
    pdf.add_page()
    content = bytes(pdf.output())

    units = parse_units("blank.pdf", content)
    assert units == []


def test_unsupported_extension_still_raises() -> None:
    with pytest.raises(ValueError):
        parse_units("image.png", b"\x89PNG\r\n")
