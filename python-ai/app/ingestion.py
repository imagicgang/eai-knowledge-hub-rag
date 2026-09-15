import ast
import csv
import html
import io
import json
import re
from pathlib import Path
from xml.etree import ElementTree

import yaml
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from .embeddings import EmbeddingProvider, HashEmbeddingProvider
from .semantic_chunking import semantic_chunk

MAX_CHUNKS = 500
CODE_SUFFIXES = {
    ".go", ".py", ".ts", ".tsx", ".js", ".jsx", ".java", ".tf", ".tfvars", ".hcl", ".cs", ".kt",
    ".cpp", ".cc", ".c", ".h", ".hpp", ".rs", ".sql", ".rb", ".php", ".swift", ".scala", ".sh",
    ".bash", ".zsh", ".ps1", ".vue", ".lua", ".pl", ".groovy", ".dart",
}
CONFIG_SUFFIXES = {".ini", ".toml", ".cfg", ".conf", ".properties", ".env"}
DOCUMENT_SUFFIXES = {".md", ".txt", ".rst", ".adoc"}
MARKUP_SUFFIXES = {".html", ".htm"}
PRESENTATION_SUFFIXES = {".pptx"}
NAMED_CODE_FILES = {"dockerfile", "makefile", "jenkinsfile", "vagrantfile"}


def parse_file(
    filename: str,
    content: bytes,
    schema: str = "",
    embedding_provider: EmbeddingProvider | None = None,
) -> list[str]:
    provider = embedding_provider or HashEmbeddingProvider()
    units = parse_units(filename, content, schema)
    return semantic_chunk(units[:MAX_CHUNKS], provider)[:MAX_CHUNKS]


def parse_units(filename: str, content: bytes, schema: str = "") -> list[str]:
    """Parse deterministic structural units without choosing a chunking strategy."""
    base_name = Path(filename).name.lower()
    suffix = Path(filename).suffix.lower()
    name_root = base_name.split(".")[0]
    if suffix == ".xlsx":
        units = _parse_xlsx_units(content, schema)
    elif suffix == ".csv":
        units = _parse_csv_units(content)
    elif suffix in {".yaml", ".yml"}:
        units = _parse_structured_units(yaml.safe_load(content.decode("utf-8-sig")))
    elif suffix == ".json":
        units = _parse_structured_units(json.loads(content.decode("utf-8-sig")))
    elif suffix == ".drawio":
        units = _parse_drawio_units(content)
    elif suffix == ".xml":
        units = _parse_drawio_units(content) or _parse_generic_xml_units(content)
    elif suffix in {".puml", ".plantuml"}:
        units = _parse_plantuml_units(content.decode("utf-8-sig"))
    elif suffix in CONFIG_SUFFIXES:
        units = _parse_config_units(content.decode("utf-8-sig"))
    elif suffix in MARKUP_SUFFIXES:
        units = _parse_html_units(content.decode("utf-8-sig"))
    elif suffix == ".docx":
        units = _parse_docx_units(content)
    elif suffix == ".pdf":
        units = _parse_pdf_units(content)
    elif suffix in PRESENTATION_SUFFIXES:
        units = _parse_pptx_units(content)
    elif suffix in CODE_SUFFIXES:
        units = _parse_code_units(content.decode("utf-8-sig"), suffix)
    elif suffix in DOCUMENT_SUFFIXES:
        units = _parse_document_units(content.decode("utf-8-sig"), suffix)
    elif not suffix and name_root in NAMED_CODE_FILES:
        units = _parse_generic_code_units(content.decode("utf-8-sig"))
    elif re.match(r"^\.env(\..+)?$", base_name):
        units = _parse_config_units(content.decode("utf-8-sig"))
    else:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")
    return units[:MAX_CHUNKS]


def _parse_xlsx_units(content: bytes, schema: str) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = [workbook[schema]] if schema and schema in workbook.sheetnames else list(workbook.worksheets)
    units: list[str] = []
    for sheet in sheets:
        rows = sheet.iter_rows(values_only=True)
        headers = [
            str(value).strip() if value is not None else f"column_{index + 1}"
            for index, value in enumerate(next(rows, []))
        ]
        for row_number, row in enumerate(rows, start=2):
            values = [
                f"{headers[index]}: {value}"
                for index, value in enumerate(row)
                if value not in (None, "")
            ]
            if values:
                units.append(f"Sheet: {sheet.title}. Row: {row_number}. " + "; ".join(values))
    return units


def _parse_csv_units(content: bytes) -> list[str]:
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    return [
        f"Row: {index}. " + "; ".join(f"{key}: {value}" for key, value in row.items() if value)
        for index, row in enumerate(reader, start=2)
        if any(row.values())
    ]


def _parse_structured_units(value: object, path: str = "root") -> list[str]:
    if isinstance(value, list):
        return [
            chunk
            for index, item in enumerate(value)
            for chunk in _parse_structured_units(item, f"{path}[{index}]")
        ]
    if isinstance(value, dict):
        scalar = [
            f"{key}: {item}" for key, item in value.items() if not isinstance(item, (dict, list))
        ]
        chunks = [f"{path}. " + "; ".join(scalar)] if scalar else []
        for key, item in value.items():
            if isinstance(item, (dict, list)):
                chunks.extend(_parse_structured_units(item, f"{path}.{key}"))
        return chunks
    return [f"{path}: {value}"]


def _parse_drawio_units(content: bytes) -> list[str]:
    root = ElementTree.fromstring(content)
    cells = {cell.get("id", ""): cell for cell in root.iter("mxCell")}
    entities = {
        cell_id: _clean_label(cell.get("value", ""))
        for cell_id, cell in cells.items()
        if cell.get("vertex") == "1" and _clean_label(cell.get("value", ""))
    }
    relations: dict[str, list[str]] = {cell_id: [] for cell_id in entities}
    standalone: list[str] = []
    for cell in cells.values():
        if cell.get("edge") != "1":
            continue
        source_id, target_id = cell.get("source", ""), cell.get("target", "")
        source, target = entities.get(source_id, source_id), entities.get(target_id, target_id)
        label = _clean_label(cell.get("value", ""))
        relation = f"Relationship: {source} -> {target}" + (f". Label: {label}" if label else "")
        if source_id in relations:
            relations[source_id].append(f"Outgoing {relation}")
        if target_id in relations:
            relations[target_id].append(f"Incoming {relation}")
        if source_id not in relations and target_id not in relations and (source or target):
            standalone.append(relation)
    units = [
        "\n".join([f"Diagram entity: {label}", *relations[cell_id]])
        for cell_id, label in entities.items()
    ]
    return units + standalone


def _xml_element_to_obj(element: ElementTree.Element) -> object:
    obj: dict[str, object] = {f"@{key}": value for key, value in element.attrib.items()}
    grouped: dict[str, list[object]] = {}
    for child in element:
        grouped.setdefault(child.tag, []).append(_xml_element_to_obj(child))
    for tag, items in grouped.items():
        obj[tag] = items if len(items) > 1 else items[0]
    text = (element.text or "").strip()
    if text:
        obj["#text"] = text
    return obj or text


def _parse_generic_xml_units(content: bytes) -> list[str]:
    root = ElementTree.fromstring(content)
    return _parse_structured_units(_xml_element_to_obj(root), path=root.tag)


def _parse_config_units(value: str) -> list[str]:
    section = "root"
    units: list[str] = []
    for raw_line in value.splitlines():
        line = raw_line.strip()
        if not line or line.startswith(("#", ";")):
            continue
        section_match = re.match(r"^\[(.+)]$", line)
        if section_match:
            section = section_match.group(1).strip()
            continue
        if "=" in line or ":" in line:
            units.append(f"Section: {section}. {line}")
    return units


def _parse_html_units(value: str) -> list[str]:
    value = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", value)
    value = re.sub(r"(?i)</(p|div|li|h[1-6]|tr|section|article)\s*>", "\n\n", value)
    value = re.sub(r"(?i)<br\s*/?>", "\n", value)
    text = html.unescape(re.sub(r"<[^>]+>", " ", value))
    text = re.sub(r"[ \t]+", " ", text)
    return [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]


def _parse_docx_units(content: bytes) -> list[str]:
    document = DocxDocument(io.BytesIO(content))
    units: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row_number, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                units.append(f"Table row {row_number}: " + "; ".join(cells))
    return units


def _parse_pdf_units(content: bytes) -> list[str]:
    """Extract text per page. PDFs without a text layer (scanned images) yield no units."""
    reader = PdfReader(io.BytesIO(content))
    units: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        for paragraph in re.split(r"\n\s*\n", text):
            paragraph = paragraph.strip()
            if paragraph:
                units.append(f"Page {page_number}: {paragraph}")
    return units


def _parse_pptx_units(content: bytes) -> list[str]:
    presentation = Presentation(io.BytesIO(content))
    units: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        lines = [
            "\n".join(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        notes = (
            slide.notes_slide.notes_text_frame.text.strip()
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame
            else ""
        )
        if not lines and not notes:
            continue
        unit = f"Slide {slide_number}: " + "\n".join(lines)
        if notes:
            unit += f"\nSpeaker notes: {notes}"
        units.append(unit.strip())
    return units


def _parse_plantuml_units(value: str) -> list[str]:
    return [
        line.strip()
        for line in value.splitlines()
        if line.strip() and not line.strip().startswith(("'", "@start", "@end"))
    ]


def _clean_label(value: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.unescape(value))).strip()


def _parse_python_units(value: str) -> list[str]:
    lines = value.splitlines()
    try:
        tree = ast.parse(value)
    except SyntaxError:
        return _parse_generic_code_units(value)
    nodes = [
        node
        for node in tree.body
        if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))
    ]
    if not nodes:
        return [value.strip()] if value.strip() else []
    starts = [max(0, node.lineno - 1) for node in nodes]
    units: list[str] = []
    for index, start in enumerate(starts):
        if index == 0:
            start = 0
        end = starts[index + 1] if index + 1 < len(starts) else len(lines)
        units.append("\n".join(lines[start:end]).strip())
    return [unit for unit in units if unit]


DECLARATION = re.compile(
    r"^\s*(?:export\s+|public\s+|private\s+|protected\s+|static\s+|async\s+)*"
    r"(?:class|interface|type|enum|function|func|def|module|package|resource|data|provider|"
    r"variable|output|impl|struct|trait)\b"
)


def _parse_generic_code_units(value: str) -> list[str]:
    lines = value.splitlines()
    starts = [index for index, line in enumerate(lines) if DECLARATION.match(line)]
    if not starts:
        paragraphs = re.split(r"\n\s*\n", value)
        return [paragraph.strip() for paragraph in paragraphs if paragraph.strip()]
    units: list[str] = []
    for index, start in enumerate(starts):
        if index == 0:
            start = 0
        end = starts[index + 1] if index + 1 < len(starts) else len(lines)
        units.append("\n".join(lines[start:end]).strip())
    return [unit for unit in units if unit]


def _parse_code_units(value: str, suffix: str) -> list[str]:
    return _parse_python_units(value) if suffix == ".py" else _parse_generic_code_units(value)


def _parse_document_units(value: str, suffix: str) -> list[str]:
    if suffix == ".md":
        starts = [match.start() for match in re.finditer(r"(?m)^#{1,6}\s+", value)]
        if starts:
            starts = ([0] if starts[0] else []) + starts
            return [
                value[start:end].strip()
                for start, end in zip(starts, starts[1:] + [len(value)], strict=True)
                if value[start:end].strip()
            ]
    return [paragraph.strip() for paragraph in re.split(r"\n\s*\n", value) if paragraph.strip()]
